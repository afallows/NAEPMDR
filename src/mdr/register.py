"""Excel master document register writer.

Layout follows normal MDR practice: identity columns first (number,
revision, status, title), then metadata, then provenance and QC. Every row
carries a hyperlink that opens the source PDF.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import theme
from .models import DocumentRecord

SHEET_REGISTER = "Document Register"
SHEET_SUMMARY = "Summary"
SHEET_QC = "QC Review"
SHEET_HISTORY = "Revision History"


@dataclass
class Column:
    key: str
    heading: str
    width: float
    align: str = "left"


COLUMNS: list[Column] = [
    Column("index", "#", 5, "center"),
    Column("link", "Open", 8, "center"),
    Column("drawing_number", "Drawing Number", 30),
    Column("revision", "Rev", 6, "center"),
    Column("status", "Status", 11, "center"),
    Column("status_description", "Latest Revision Description", 52),
    Column("revision_date", "Rev Date", 12, "center"),
    Column("title", "Title", 56),
    Column("facility", "Facility / Location", 34),
    Column("discipline", "Discipline", 16),
    Column("sheet", "Sheet", 8, "center"),
    Column("page_count", "Sheets", 8, "center"),
    Column("project_number", "Project No.", 14, "center"),
    Column("scale", "Scale", 10, "center"),
    Column("drawing_date", "Drawing Date", 13, "center"),
    Column("drawn_by", "Drawn / Design", 15),
    Column("drawing_check", "Drawing Check", 15),
    Column("design_engineer", "Design Engineer", 16),
    Column("design_check", "Design Check", 14),
    Column("project_approval", "Project Approval", 16),
    Column("file_name", "File Name", 40),
    Column("folder", "Folder", 32),
    Column("file_type", "Type", 7, "center"),
    Column("companion_dwg", "DWG Available", 26),
    Column("source", "Source", 18),
    Column("confidence", "Confidence", 11, "center"),
    Column("qc_flags", "QC Flags", 60),
    Column("file_size_kb", "Size (KB)", 11, "center"),
    Column("file_modified", "File Modified", 17, "center"),
    Column("file_path", "Full Path", 60),
]

TITLE_ROW = 1
SUBTITLE_ROW = 2
HEADER_ROW = 4
FIRST_DATA_ROW = 5

_THIN = Side(style="thin", color=theme.GREY_400)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _fill(colour: str) -> PatternFill:
    return PatternFill("solid", fgColor=colour, bgColor=colour)


def _cell_value(record: DocumentRecord, key: str, index: int):
    if key == "index":
        return index
    if key == "link":
        return "Open"
    if key == "sheet":
        return record.page
    if key == "qc_flags":
        return record.qc_text
    if key == "confidence":
        return round(record.confidence / 100.0, 3) if record.confidence else None
    if key == "file_size_kb":
        return round(record.file_size_bytes / 1024.0, 1) if record.file_size_bytes else None
    if key == "companion_dwg":
        return record.companion_dwg or ""
    return getattr(record, key, "")


class RegisterWriter:
    """Builds the workbook."""

    def __init__(self, records: list[DocumentRecord], root: str,
                 include_full_path: bool = True):
        self.records = records
        self.root = root
        self.columns = [
            c for c in COLUMNS
            if include_full_path or c.key != "file_path"
        ]

    # -- shared styling --------------------------------------------------

    def _banner(self, sheet: Worksheet, title: str, subtitle: str, span: int) -> None:
        """Enbridge-branded banner across the top of a sheet."""
        sheet.merge_cells(start_row=TITLE_ROW, start_column=1,
                          end_row=TITLE_ROW, end_column=max(1, span))
        cell = sheet.cell(row=TITLE_ROW, column=1, value=title)
        cell.font = Font(name=theme.FONT_NAME, size=theme.TITLE_SIZE, bold=True,
                         color=theme.ENBRIDGE_DARK_GREY)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[TITLE_ROW].height = 30

        sheet.merge_cells(start_row=SUBTITLE_ROW, start_column=1,
                          end_row=SUBTITLE_ROW, end_column=max(1, span))
        sub = sheet.cell(row=SUBTITLE_ROW, column=1, value=subtitle)
        sub.font = Font(name=theme.FONT_NAME, size=theme.SUBTITLE_SIZE,
                        color=theme.GREY_600)
        sub.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[SUBTITLE_ROW].height = 16

        # Gold rule under the banner.
        for col in range(1, max(2, span + 1)):
            rule = sheet.cell(row=SUBTITLE_ROW + 1, column=col)
            rule.fill = _fill(theme.ENBRIDGE_GOLD)
        sheet.row_dimensions[SUBTITLE_ROW + 1].height = 4

    def _header_row(self, sheet: Worksheet, row: int, headings: list[str],
                    widths: list[float] | None = None) -> None:
        for i, heading in enumerate(headings, start=1):
            cell = sheet.cell(row=row, column=i, value=heading)
            cell.font = Font(name=theme.FONT_NAME, size=theme.HEADER_SIZE,
                             bold=True, color=theme.WHITE)
            cell.fill = _fill(theme.ENBRIDGE_DARK_GREY)
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
            cell.border = _BORDER
        sheet.row_dimensions[row].height = 30
        if widths:
            for i, width in enumerate(widths, start=1):
                sheet.column_dimensions[get_column_letter(i)].width = width

    # -- register sheet ---------------------------------------------------

    def _write_register(self, sheet: Worksheet) -> None:
        span = len(self.columns)
        self._banner(
            sheet,
            "Master Document Register",
            f"Source folder: {self.root}    |    Generated "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M')}    |    "
            f"{len(self.records)} sheet(s)",
            span,
        )
        self._header_row(sheet, HEADER_ROW, [c.heading for c in self.columns],
                         [c.width for c in self.columns])

        band = _fill(theme.GREY_100)
        for offset, record in enumerate(self.records):
            row = FIRST_DATA_ROW + offset
            striped = offset % 2 == 1
            for col_index, column in enumerate(self.columns, start=1):
                cell = sheet.cell(row=row, column=col_index,
                                  value=_cell_value(record, column.key, offset + 1))
                cell.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE,
                                 color=theme.ENBRIDGE_BLACK)
                cell.alignment = Alignment(
                    horizontal=column.align,
                    vertical="top",
                    wrap_text=column.key in ("title", "status_description", "qc_flags",
                                             "facility", "folder", "file_path"),
                )
                cell.border = _BORDER
                if striped:
                    cell.fill = band

                if column.key == "link":
                    self._style_link(cell, record)
                elif column.key == "status":
                    cell.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE,
                                     bold=True, color=theme.status_colour(record.status))
                elif column.key == "drawing_number":
                    cell.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE,
                                     bold=True, color=theme.ENBRIDGE_BLACK)
                elif column.key == "confidence":
                    cell.number_format = "0%"
                elif column.key == "qc_flags" and record.qc_flags:
                    cell.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE,
                                     color=theme.QC_WARN)
                elif column.key == "file_size_kb":
                    cell.number_format = "#,##0.0"

            if record.has_error:
                for col_index in range(1, span + 1):
                    sheet.cell(row=row, column=col_index).fill = _fill("FBE3DF")

        last_row = FIRST_DATA_ROW + len(self.records) - 1
        if self.records:
            sheet.auto_filter.ref = (
                f"A{HEADER_ROW}:{get_column_letter(span)}{last_row}"
            )
            self._conditional_formats(sheet, last_row)
        sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=4)
        sheet.sheet_view.showGridLines = False

    def _style_link(self, cell, record: DocumentRecord) -> None:
        """Turn the 'Open' cell into a hyperlink to the source document."""
        target = record.link_target or record.file_path
        if not target:
            cell.value = ""
            return
        # Excel wants forward slashes and a page anchor for multi-sheet PDFs.
        location = target.replace("\\", "/")
        if record.page_count > 1 and record.page > 1:
            location = f"{location}#page={record.page}"
        cell.hyperlink = location
        cell.value = "Open"
        cell.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE,
                         color="0B5FA5", underline="single", bold=True)

    def _conditional_formats(self, sheet: Worksheet, last_row: int) -> None:
        """Colour-code confidence so weak reads stand out."""
        keys = [c.key for c in self.columns]
        if "confidence" not in keys:
            return
        letter = get_column_letter(keys.index("confidence") + 1)
        span = f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}"
        sheet.conditional_formatting.add(span, CellIsRule(
            operator="lessThan", formula=["0.6"],
            font=Font(color=theme.QC_ERROR, bold=True),
            fill=_fill("FBE3DF")))
        sheet.conditional_formatting.add(span, CellIsRule(
            operator="between", formula=["0.6", "0.8"],
            font=Font(color=theme.QC_WARN, bold=True)))
        sheet.conditional_formatting.add(span, CellIsRule(
            operator="greaterThan", formula=["0.8"],
            font=Font(color=theme.QC_OK, bold=True)))

    # -- summary ----------------------------------------------------------

    def _write_summary(self, sheet: Worksheet) -> None:
        total = len(self.records)
        flagged = sum(1 for r in self.records if r.qc_flags)
        errored = sum(1 for r in self.records if r.has_error)
        with_dwg = sum(1 for r in self.records if r.companion_dwg)
        confidences = [r.confidence for r in self.records if r.confidence]
        mean_conf = sum(confidences) / len(confidences) if confidences else 0

        self._banner(sheet, "Register Summary",
                     f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}", 4)

        row = HEADER_ROW
        self._header_row(sheet, row, ["Metric", "Value", "", ""], [40, 18, 4, 40])
        metrics = [
            ("Sheets registered", total),
            ("Unique drawing numbers", len({r.drawing_number for r in self.records if r.drawing_number})),
            ("Source files", len({r.file_path for r in self.records})),
            ("Sheets with a companion DWG", with_dwg),
            ("Rows with QC flags", flagged),
            ("Rows that failed to extract", errored),
            ("Mean extraction confidence", f"{mean_conf:.0f}%"),
        ]
        row += 1
        for name, value in metrics:
            self._kv(sheet, row, name, value)
            row += 1

        row += 1
        self._header_row(sheet, row, ["Status", "Sheets", "", ""], [40, 18, 4, 40])
        row += 1
        for status, count in Counter(
            (r.status or "(not determined)") for r in self.records
        ).most_common():
            self._kv(sheet, row, status, count,
                     colour=theme.status_colour(status if status != "(not determined)" else None))
            row += 1

        row += 1
        self._header_row(sheet, row, ["Discipline", "Sheets", "", ""], [40, 18, 4, 40])
        row += 1
        for discipline, count in Counter(
            (r.discipline or "(unclassified)") for r in self.records
        ).most_common():
            self._kv(sheet, row, discipline, count)
            row += 1

        sheet.sheet_view.showGridLines = False

    def _kv(self, sheet: Worksheet, row: int, name, value, colour: str | None = None) -> None:
        left = sheet.cell(row=row, column=1, value=name)
        left.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE,
                         color=colour or theme.ENBRIDGE_BLACK,
                         bold=colour is not None)
        left.border = _BORDER
        right = sheet.cell(row=row, column=2, value=value)
        right.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE, bold=True)
        right.alignment = Alignment(horizontal="center")
        right.border = _BORDER

    # -- QC ---------------------------------------------------------------

    def _write_qc(self, sheet: Worksheet) -> None:
        flagged = [r for r in self.records if r.qc_flags or r.has_error]
        self._banner(
            sheet, "QC Review",
            "Rows needing a human check. Everything here is also in the main "
            "register - this sheet is the worklist.",
            6,
        )
        headings = ["#", "Open", "Drawing Number", "Rev", "File Name", "Issue"]
        self._header_row(sheet, HEADER_ROW, headings, [5, 8, 30, 6, 42, 90])
        for offset, record in enumerate(flagged):
            row = FIRST_DATA_ROW + offset
            issue = record.error or record.qc_text
            values = [offset + 1, "Open", record.drawing_number, record.revision,
                      record.file_name, issue]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=col, value=value)
                cell.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE)
                cell.alignment = Alignment(vertical="top", wrap_text=col == 6)
                cell.border = _BORDER
                if offset % 2:
                    cell.fill = _fill(theme.GREY_100)
            self._style_link(sheet.cell(row=row, column=2), record)
            sheet.cell(row=row, column=6).font = Font(
                name=theme.FONT_NAME, size=theme.BODY_SIZE,
                color=theme.QC_ERROR if record.has_error else theme.QC_WARN)
        if flagged:
            sheet.auto_filter.ref = f"A{HEADER_ROW}:F{FIRST_DATA_ROW + len(flagged) - 1}"
        else:
            note = sheet.cell(row=FIRST_DATA_ROW, column=1,
                              value="No issues flagged - every sheet extracted cleanly.")
            note.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE,
                             color=theme.QC_OK, bold=True)
        sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=1)
        sheet.sheet_view.showGridLines = False

    # -- revision history --------------------------------------------------

    def _write_history(self, sheet: Worksheet) -> None:
        self._banner(sheet, "Revision History",
                     "Every revision row read from each drawing's revision table.", 6)
        headings = ["Drawing Number", "Sheet", "Rev", "Date", "Description", "File Name"]
        self._header_row(sheet, HEADER_ROW, headings, [30, 8, 6, 13, 78, 40])
        row = FIRST_DATA_ROW
        shaded = False
        for record in self.records:
            if not record.revision_history:
                continue
            for entry in record.revision_history:
                values = [record.drawing_number, record.page, entry.revision,
                          entry.date, entry.description, record.file_name]
                for col, value in enumerate(values, start=1):
                    cell = sheet.cell(row=row, column=col, value=value)
                    cell.font = Font(name=theme.FONT_NAME, size=theme.BODY_SIZE)
                    cell.alignment = Alignment(vertical="top", wrap_text=col == 5)
                    cell.border = _BORDER
                    if shaded:
                        cell.fill = _fill(theme.GREY_100)
                row += 1
            shaded = not shaded
        if row > FIRST_DATA_ROW:
            sheet.auto_filter.ref = f"A{HEADER_ROW}:F{row - 1}"
        sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=1)
        sheet.sheet_view.showGridLines = False

    # -- entry point -------------------------------------------------------

    def write(self, output_path: str) -> str:
        workbook = Workbook()
        register = workbook.active
        register.title = SHEET_REGISTER
        self._write_register(register)
        self._write_summary(workbook.create_sheet(SHEET_SUMMARY))
        self._write_qc(workbook.create_sheet(SHEET_QC))
        self._write_history(workbook.create_sheet(SHEET_HISTORY))

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(path))
        return str(path)


def write_register(records: list[DocumentRecord], root: str, output_path: str,
                   include_full_path: bool = True) -> str:
    return RegisterWriter(records, root, include_full_path).write(output_path)
