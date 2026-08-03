"""Tests for the Excel register writer."""

import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from mdr import theme  # noqa: E402
from mdr.models import DocumentRecord, RevisionEntry  # noqa: E402
from mdr.register import FIRST_DATA_ROW, HEADER_ROW, write_register  # noqa: E402


def _record(**kw) -> DocumentRecord:
    base = dict(
        drawing_number="PP-1-94000-5530-01", revision="B", status="IFB",
        status_description="RE-ISSUED FOR BID (30006386 NORTH AITKEN EXPANSION)",
        title="PANEL LAYOUT AND BILL OF MATERIAL",
        facility="PP-1 AITKEN CREEK GAS STORAGE",
        file_path="/root/EHT/PP-1-94000-5530-01_RB.pdf",
        relative_path="EHT/PP-1-94000-5530-01_RB.pdf",
        link_target="EHT/PP-1-94000-5530-01_RB.pdf",
        file_name="PP-1-94000-5530-01_RB.pdf", folder="EHT", file_type="PDF",
        page=1, page_count=1, confidence=88.3,
        revision_history=[RevisionEntry("B", "2026-05-08", "RE-ISSUED FOR BID")],
    )
    base.update(kw)
    return DocumentRecord(**base)


@pytest.fixture
def book(tmp_path):
    def build(records):
        path = tmp_path / "register.xlsx"
        write_register(records, "/root", str(path))
        return load_workbook(str(path))
    return build


class TestStructure:
    def test_expected_sheets(self, book):
        wb = book([_record()])
        assert wb.sheetnames == [
            "Document Register", "Summary", "QC Review", "Revision History"]

    def test_header_and_freeze(self, book):
        ws = book([_record()])["Document Register"]
        assert ws.cell(HEADER_ROW, 3).value == "Drawing Number"
        assert ws.cell(HEADER_ROW, 5).value == "Status"
        assert ws.freeze_panes == "D5"
        assert ws.auto_filter.ref is not None

    def test_one_row_per_record(self, book):
        records = [_record(), _record(drawing_number="X-2", page=1)]
        ws = book(records)["Document Register"]
        assert ws.max_row == FIRST_DATA_ROW + len(records) - 1


class TestBranding:
    def test_gold_rule_and_grey_header(self, book):
        ws = book([_record()])["Document Register"]
        assert ws.cell(3, 1).fill.fgColor.rgb.endswith(theme.ENBRIDGE_GOLD)
        assert ws.cell(HEADER_ROW, 1).fill.fgColor.rgb.endswith(
            theme.ENBRIDGE_DARK_GREY)

    def test_status_is_colour_coded(self, book):
        ws = book([_record(status="IFC"), _record(status="AS-BUILT")])["Document Register"]
        assert ws.cell(FIRST_DATA_ROW, 5).font.color.rgb.endswith(
            theme.STATUS_COLOURS["IFC"])
        assert ws.cell(FIRST_DATA_ROW + 1, 5).font.color.rgb.endswith(
            theme.STATUS_COLOURS["AS-BUILT"])


class TestHyperlinks:
    def test_single_sheet_links_to_the_file(self, book):
        ws = book([_record()])["Document Register"]
        assert ws.cell(FIRST_DATA_ROW, 2).hyperlink.target == \
            "EHT/PP-1-94000-5530-01_RB.pdf"

    def test_multi_sheet_links_carry_a_page_anchor(self, book):
        ws = book([_record(page=4, page_count=9)])["Document Register"]
        assert ws.cell(FIRST_DATA_ROW, 2).hyperlink.target.endswith("#page=4")

    def test_backslashes_become_forward_slashes(self, book):
        ws = book([_record(link_target=r"EHT\sub\file.pdf")])["Document Register"]
        assert "\\" not in ws.cell(FIRST_DATA_ROW, 2).hyperlink.target

    def test_hash_in_path_suppresses_the_page_anchor(self, book):
        # '#' is the fragment separator; appending an anchor would produce a
        # link that resolves to a truncated path and fails to open.
        ws = book([_record(link_target="EHT/rev#2/file.pdf", page=3,
                           page_count=5)])["Document Register"]
        assert ws.cell(FIRST_DATA_ROW, 2).hyperlink.target == "EHT/rev#2/file.pdf"

    def test_missing_target_leaves_the_cell_blank(self, book):
        ws = book([_record(link_target="", file_path="")])["Document Register"]
        assert ws.cell(FIRST_DATA_ROW, 2).hyperlink is None


class TestQcSheet:
    def test_only_flagged_rows_appear(self, book):
        clean = _record()
        flagged = _record(drawing_number="X-9",
                          qc_flags=["Revision mismatch - title block 'B' vs filename 'A'"])
        ws = book([clean, flagged])["QC Review"]
        assert ws.cell(FIRST_DATA_ROW, 3).value == "X-9"
        assert ws.cell(FIRST_DATA_ROW + 1, 3).value is None

    def test_errors_are_reported_over_flags(self, book):
        ws = book([_record(error="Cannot open PDF: damaged",
                           qc_flags=["something else"])])["QC Review"]
        assert "damaged" in ws.cell(FIRST_DATA_ROW, 6).value

    def test_clean_scan_says_so(self, book):
        ws = book([_record()])["QC Review"]
        assert "No issues" in ws.cell(FIRST_DATA_ROW, 1).value


class TestHistorySheet:
    def test_every_revision_row_is_written(self, book):
        record = _record(revision_history=[
            RevisionEntry("A", "2026-04-14", "ISSUED FOR BID"),
            RevisionEntry("B", "2026-05-08", "RE-ISSUED FOR BID"),
        ])
        ws = book([record])["Revision History"]
        assert ws.cell(FIRST_DATA_ROW, 3).value == "A"
        assert ws.cell(FIRST_DATA_ROW + 1, 3).value == "B"


class TestSummarySheet:
    def test_counts_reflect_the_records(self, book):
        ws = book([_record(), _record(status="IFC"),
                   _record(qc_flags=["x"])])["Summary"]
        values = {ws.cell(r, 1).value: ws.cell(r, 2).value
                  for r in range(HEADER_ROW, ws.max_row + 1)}
        assert values["Sheets registered"] == 3
        assert values["Rows with QC flags"] == 1
